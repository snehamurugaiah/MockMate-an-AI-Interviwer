
from flask import Flask, request, render_template, jsonify, redirect, url_for, session, flash
import os
import logging
import psycopg2
from transformers import pipeline
from random import randint
import pytesseract
from PyPDF2 import PdfReader
from docx import Document
import re
from flask import send_file
from fpdf import FPDF
from io import BytesIO
app = Flask(__name__)
app.secret_key = 'your_secret_key'  # Secret key for session management
recordings_path = 'static/recordings'

# Logging Configuration
logging.basicConfig(level=logging.INFO)

# Database Configuration
DB_CONFIG = {
    "dbname": "postgres",
    "user": "postgres",
    "password": "dhivya@2005",
    "host": "localhost",
    "port": "5432",
}

def get_connection():
    return psycopg2.connect(**DB_CONFIG)

def fetch_questions(table_name, limit=10):  # Fetch 10 questions instead of 5
    conn = get_connection()
    cursor = conn.cursor()
    query = f"SELECT job_role, question, difficulty, source_api FROM {table_name} ORDER BY RANDOM() LIMIT {limit}"
    cursor.execute(query)
    data = cursor.fetchall()
    conn.close()
    return data

# Hugging Face Pipelines
#asr = pipeline("automatic-speech-recognition", model="facebook/wav2vec2-large-960h", framework="pt")
asr = pipeline("automatic-speech-recognition", model="openai/whisper-tiny.en")  # or another TF-compatible model

emotion_recog = pipeline("text-classification", model="j-hartmann/emotion-english-distilroberta-base")

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username']
        password = request.form['password']

        # Hardcoded admin login (replace this with a DB check in production)
        if username == 'admin' and password == 'password123':
            session.clear()  # Clear any existing session data
            session['user'] = username
            session['scores'] = []
            session['confidence_scores'] = []
            session['resume_text'] = ''
            flash("Login successful!", "success")
            return redirect(url_for('upload_resume'))
        else:
            flash("Invalid credentials, please try again.", "danger")

    return render_template('login.html')

def extract_text_from_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        text += page.extract_text()
    return text

def extract_text_from_docx(file_path):
    doc = Document(file_path)
    return "\n".join([para.text for para in doc.paragraphs])


def extract_features_from_text(text):
    features = {
        "skills": [],
        "experience": [],
        "certifications": [],
        "projects": []
    }

    # Sample patterns (these can be improved or replaced with NLP later)
    skill_keywords = ["python", "java", "sql", "html", "css", "machine learning", "deep learning"]
    cert_keywords = ["certified", "certificate", "certification"]
    exp_keywords = ["experience", "worked at", "intern", "internship", "project"]
    project_keywords = ["project", "developed", "built", "created"]

    lower_text = text.lower()

    features["skills"] = [kw for kw in skill_keywords if kw in lower_text]
    features["certifications"] = re.findall(r"(certified[^.\n]*)", text, re.IGNORECASE)
    features["experience"] = re.findall(r"(experience[^.\n]*)", text, re.IGNORECASE)
    features["projects"] = re.findall(r"(project[^.\n]*)", text, re.IGNORECASE)

    return features

@app.route('/upload_resume', methods=['GET', 'POST'])
def upload_resume():
    if 'user' not in session:
        flash("Login required", "warning")
        return redirect(url_for('login'))

    extracted_features = {}

    if request.method == 'POST':
        resume = request.files['resume']
        filename = os.path.join('static/uploads', resume.filename)
        os.makedirs('static/uploads', exist_ok=True)
        resume.save(filename)

        # Extract text
        if filename.endswith('.pdf'):
            extracted_text = extract_text_from_pdf(filename)
        elif filename.endswith('.docx'):
            extracted_text = extract_text_from_docx(filename)
        else:
            flash("Unsupported file format. Use PDF or DOCX.", "danger")
            return redirect(url_for('upload_resume'))

        # Extract features from resume text
        extracted_features = extract_features_from_text(extracted_text)

        session['resume_text'] = extracted_text  # Optional: for future use

        flash("Resume uploaded and features extracted!", "success")

    return render_template('upload_resume.html', features=extracted_features)


@app.route('/logout')
def logout():
    session.clear()
    flash("You have been logged out.", "info")
    return redirect(url_for('login'))


@app.route('/')
def index():
    print("Session Data:", session)  # Debugging
    if 'user' not in session:
        flash("You need to log in first.", "warning")
        return redirect(url_for('login'))

    ds_questions = fetch_questions("job_questions", limit=10)  # Now fetches 10 questions
    hr_questions = fetch_questions("hr_questions", limit=10)  # Now fetches 10 questions
    return render_template('index.html', ds_questions=ds_questions, hr_questions=hr_questions)

@app.route('/submit_answer', methods=['POST'])
def submit_answer():
    if 'user' not in session:
        return jsonify({"error": "Unauthorized access"}), 401

    if 'audio_data' not in request.files:
        return jsonify({"error": "No audio file uploaded"}), 400

    audio_file = request.files['audio_data']
    question_id = request.form.get('question_id')

    # Generate a random confidence score between 0 and 100
    confidence_score = randint(0, 100)

    # Store score in session
    if 'scores' not in session:
        session['scores'] = []
    if 'confidence_scores' not in session:
        session['confidence_scores'] = []

    session['scores'].append(confidence_score)
    session['confidence_scores'].append(confidence_score)

    session.modified = True  # Important to ensure session is updated

    return jsonify({"confidence_score": confidence_score})
@app.route('/report')
def report():
    if 'scores' not in session or not session['scores']:
        flash("No scores available to generate report.", "warning")
        return redirect(url_for('index'))

    scores = session['scores']
    avg_score = sum(scores) / len(scores)

    # Command/suggestion based on average score
    if avg_score >= 80:
        suggestion = "Excellent performance! Ready for real interviews."
    elif avg_score >= 60:
        suggestion = "Good job! A little more practice will help."
    else:
        suggestion = "Needs improvement. Focus on communication and confidence."

    return render_template('report.html', scores=scores, avg_score=avg_score, suggestion=suggestion)

@app.route('/download_report')
def download_report():
    if 'confidence_scores' not in session:
        flash("No report data found.", "warning")
        return redirect(url_for('index'))

    scores = session['confidence_scores']
    avg_score = sum(scores) / len(scores)
    suggestion = session.get('suggestion', 'No suggestion available.')

    # Generate PDF in memory
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    pdf.cell(200, 10, txt="Mock Interview Report", ln=True, align='C')
    pdf.ln(10)

    for i, score in enumerate(scores, start=1):
        pdf.cell(200, 10, txt=f"Question {i}: {score}%", ln=True)

    pdf.ln(5)
    pdf.cell(200, 10, txt=f"Average Score: {avg_score:.2f}%", ln=True)
    pdf.ln(5)
    pdf.multi_cell(0, 10, txt=f"Suggestion:\n{suggestion}")

    # Save PDF to BytesIO buffer
    buffer = BytesIO()
    pdf.output(buffer)
    buffer.seek(0)

    return send_file(
        buffer,
        as_attachment=True,
        download_name="interview_report.pdf",
        mimetype='application/pdf'
    )




if __name__ == '__main__':
    if not os.path.exists(recordings_path):
        os.makedirs(recordings_path)
    app.run(debug=True)
